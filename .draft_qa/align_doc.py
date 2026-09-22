from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root = Path(__file__).resolve().parent.parent
doc = Document(root / 'Ambulance Dispatch Design.docx')
for p in list(doc.paragraphs):
    p._element.getparent().remove(p._element)
for name in ['Normal', 'Title', 'Heading 1']:
    style = doc.styles[name]
    style.font.name = 'Times New Roman'
    style.font.size = Pt(12)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.bold = name != 'Normal'
    style.font.italic = False
    fonts = style.element.rPr.rFonts
    for attr in list(fonts.attrib):
        if 'theme' in attr.lower(): del fonts.attrib[attr]
    for attr in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
        fonts.set(qn('w:' + attr), 'Times New Roman')
    for tag in ['spacing', 'kern']:
        for element in list(style.element.rPr.findall(qn('w:' + tag))):
            style.element.rPr.remove(element)
    fmt = style.paragraph_format
    fmt.line_spacing = 2
    fmt.space_before = fmt.space_after = Pt(0)
    fmt.first_line_indent = Inches(.5 if name == 'Normal' else 0)
    fmt.alignment = WD_ALIGN_PARAGRAPH.LEFT if name == 'Normal' else WD_ALIGN_PARAGRAPH.CENTER
    fmt.widow_control = True
for sec in doc.sections:
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.different_first_page_header_footer = False
    sec.header_distance = Inches(.5)
    for part in [sec.header, sec.footer]:
        for p in part.paragraphs: p.clear()
    p = sec.header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.first_line_indent = Inches(0)
    p.add_run().font.name = 'Times New Roman'
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); p._p.append(field)
for _ in range(3): doc.add_paragraph()
doc.add_paragraph('Ambulance Dispatch Application', 'Title')
doc.add_paragraph()
for text in ['Lavonte Wynn', 'Department of Technology, Western Governors University',
             'D795: Applied Algorithms and Reasoning', 'Dr. Thielfoldt', 'September 21, 2026']:
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Inches(0)
doc.add_page_break()
doc.add_paragraph('Ambulance Dispatch Design', 'Title')
sections = [
    ('Overall Operation', 'I am building two Python ambulance dispatch prototypes. Both currently load the simulation data and preview the first 10 calls in dispatch order. I have added the dispatch structure, but I still need to choose and implement the two routing algorithms. Once routing is ready, each prototype will compare routes from every ambulance’s staging location, include traffic delays, and select the ambulance with the lowest travel time. It will append the call ID, type, location, selected ambulance, route, and travel time to /var/log/ambulance_call_log.csv, then reset the ambulance to its staging location. A timer will measure the total time spent calculating routes.'),
    ('Simulation Files', 'Both files use Python’s csv module to read the four CSVs from the data folder. The ambulance.csv file supplies ambulance numbers and staging locations, call_priority.csv supplies priorities by call type, and calls.csv supplies call IDs, locations, and types. The location_network.csv file becomes a directed graph that stores distance and travel time plus traffic delay for each connection. The code checks required fields, numeric values, duplicate IDs, call types, and locations before processing calls.'),
    ('Call Priorities', 'Both prototypes process priority 1 calls first, followed by priority 2 and then priority 3. The code looks up each call’s priority by its call type and sorts calls by priority and original row position. Since calls.csv has no arrival timestamps, I use file order as the order received. This keeps calls with the same priority in arrival order and gives both prototypes the same sequence for comparison.')
]
for heading, text in sections:
    doc.add_paragraph(heading, 'Heading 1')
    doc.add_paragraph(text)
doc.save(root / 'Ambulance Dispatch Design.docx')
print('Saved title page and three body paragraphs.')
