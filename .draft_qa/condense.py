from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

root = Path(__file__).resolve().parent.parent
doc = Document(r'C:/Users/jorda/OneDrive/Desktop/Ambulance Dispatch Design.docx')
for p in list(doc.paragraphs):
    p._element.getparent().remove(p._element)
normal = doc.styles['Normal']
normal.font.name = 'Times New Roman'
normal.font.size = Pt(12)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.font.bold = False
normal.font.italic = False
fonts = normal.element.rPr.rFonts
for attr in list(fonts.attrib):
    if 'theme' in attr.lower(): del fonts.attrib[attr]
normal.paragraph_format.line_spacing = 2
normal.paragraph_format.space_before = normal.paragraph_format.space_after = Pt(0)
normal.paragraph_format.first_line_indent = Inches(.5)
normal.paragraph_format.keep_with_next = False
normal.paragraph_format.widow_control = True
for section in doc.sections:
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(1)
    section.page_width, section.page_height = Inches(8.5), Inches(11)

paragraphs = [
    ('Overall operation. ', 'I plan to build a Python application that dispatches ambulances to emergency calls using the simulation data. For each call, it will compare the fastest route from every ambulance’s staging location, including traffic delays, and select the ambulance with the lowest travel time. It will append the call ID, call type, call location, selected ambulance, route, and travel time to /var/log/ambulance_call_log.csv as comma-separated name-value pairs. After logging the dispatch, it will reset the ambulance to its staging location and process the next call. Both prototypes will follow this design, with different routing algorithms that I will choose later. Embedded counters will measure the total execution time spent finding routes.'),
    ('Simulation files. ', 'I will use Python’s csv module to read the four files. The ambulance.csv file provides ambulance numbers and staging locations, call_priority.csv connects call types to priorities, and calls.csv provides call IDs, locations, and types. The location_network.csv file provides connections between locations, distances, travel times, and traffic delays. I will store the network as a directed graph because travel values can differ by direction. Each connection’s cost will be its travel time plus traffic delay. Before dispatching, the application will check required fields, convert numeric values, and verify that call types and locations match the supporting data.'),
    ('Call priorities. ', 'I will process priority 1 calls first, followed by priority 2 and then priority 3. The application will look up each call’s priority using its call type and keep calls with the same priority in the order received. Since calls.csv has no arrival timestamps, I will use its row order as the arrival order. For example, a later priority 1 call will be handled before an earlier priority 2 call, but it will stay behind earlier priority 1 calls. Both prototypes will use the same call order.')
]
for label, text in paragraphs:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.add_run(label).bold = True
    p.add_run(text)
doc.save(root / 'Ambulance Dispatch Design.docx')
assert len(doc.paragraphs) == 3
assert 'WGU' not in '\n'.join(p.text for p in doc.paragraphs)
print('Saved three paragraphs:', sum(len(p.text.split()) for p in doc.paragraphs), 'words')
